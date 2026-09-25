"""
结果文件写出模块
直接加载附件5 的官方模板并填充，保证行列标签与格式与模板完全一致。
（模板时段标签比附件标签滞后一格，本模块按"位置一一对应"填充，见论文假设说明）
"""
import numpy as np
import openpyxl
from pathlib import Path
from cumcm_core import DATA_DIR, N_T, DT

TPL = DATA_DIR.parent.parent / "2026国赛_C题" / "附件" / "附件5"
if not TPL.exists():
    TPL = DATA_DIR / "附件5"

BLOCK_EDGES = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]
ND = 4          # 保留小数位


def _r(x):
    v = round(float(x), ND)
    return 0.0 if v == 0 else v          # 避免 -0.0


# ------------------------------------------------------------ result1
def write_result1(path, purch_kwh, u, v, E, E_init):
    wb = openpyxl.load_workbook(TPL / "result1.xlsx")
    ws = wb["计划购电量"]
    for i in range(N_T):
        ws.cell(row=2 + i, column=2).value = _r(purch_kwh[i])

    ws = wb["充放电量"]
    for bi, (a, b) in enumerate(BLOCK_EDGES):
        r = 2 + bi
        ws.cell(row=r, column=2).value = _r(u[a:b].sum() * DT)
        ws.cell(row=r, column=3).value = _r(v[a:b].sum() * DT)
    ws.cell(row=2, column=5).value = _r(E_init)
    ws.cell(row=3, column=5).value = _r(E[-1])
    wb.save(path)


# ------------------------------------------------------------ result2 / 4-2
def write_result2_like(path, dates, purch, u, v, E_start_arr, E_end_arr,
                       emergency=None, tpl_name="result2.xlsx", fee_day=None):
    """
    purch: (D,144) 计划购电量 kWh
    u,v  : (D,144) 充/放电功率 kW
    E_start_arr, E_end_arr: (D,) 每日 0:00 / 24:00 储电量
    fee_day: (D,) 每日购电费
    """
    wb = openpyxl.load_workbook(TPL / tpl_name)
    ws = wb["计划购电量"]
    D = len(dates)
    for i in range(D):
        r = 2 + i
        for k in range(N_T):
            ws.cell(row=r, column=2 + k).value = _r(purch[i, k])
        ws.cell(row=r, column=2 + N_T).value = _r(purch[i].sum())
        if fee_day is not None:
            ws.cell(row=r, column=2 + N_T + 1).value = _r(fee_day[i])

    ws = wb["充放电量"]
    for i in range(D):
        base = 2 + i * 6
        for bi, (a, b) in enumerate(BLOCK_EDGES):
            ws.cell(row=base + bi, column=3).value = _r(u[i, a:b].sum() * DT)
            ws.cell(row=base + bi, column=4).value = _r(v[i, a:b].sum() * DT)
        ws.cell(row=base, column=6).value = _r(E_start_arr[i])
        ws.cell(row=base + 1, column=6).value = _r(E_end_arr[i])

    ws = wb["紧急购电量"]
    if emergency:
        r = 2
        for day, slot, kwh in emergency:
            ws.cell(row=r, column=1).value = dates[day]
            ws.cell(row=r, column=2).value = slot
            ws.cell(row=r, column=3).value = _r(kwh)
            r += 1
    wb.save(path)


# ------------------------------------------------------------ result3 / 4-3
def write_result3_like(path, dates, plan_kwh, adj_kwh, u, v, E_start_arr, E_end_arr,
                       emergency=None, tpl_name="result3.xlsx",
                       fee_plan=None, fee_adj=None):
    """
    plan_kwh / adj_kwh: (D,144) 购电量【kWh】（注意：不是功率 kW）
    u, v              : (D,144) 充/放电【功率 kW】，内部乘 Δt 转电量
    """
    wb = openpyxl.load_workbook(TPL / tpl_name)
    D = len(dates)
    for sheet, arr, fees in (("计划购电量", plan_kwh, fee_plan),
                             ("调整购电量", adj_kwh, fee_adj)):
        ws = wb[sheet]
        for i in range(D):
            r = 2 + i
            for k in range(N_T):
                ws.cell(row=r, column=2 + k).value = _r(arr[i, k])
            ws.cell(row=r, column=2 + N_T).value = _r(arr[i].sum())
            if fees is not None:
                ws.cell(row=r, column=2 + N_T + 1).value = _r(fees[i])

    ws = wb["充放电量"]
    for i in range(D):
        base = 2 + i * 6
        for bi, (a, b) in enumerate(BLOCK_EDGES):
            ws.cell(row=base + bi, column=3).value = _r(u[i, a:b].sum() * DT)
            ws.cell(row=base + bi, column=4).value = _r(v[i, a:b].sum() * DT)
        ws.cell(row=base, column=6).value = _r(E_start_arr[i])
        ws.cell(row=base + 1, column=6).value = _r(E_end_arr[i])

    ws = wb["紧急购电量"]
    if emergency:
        r = 2
        for day, slot, kwh in emergency:
            ws.cell(row=r, column=1).value = dates[day]
            ws.cell(row=r, column=2).value = slot
            ws.cell(row=r, column=3).value = _r(kwh)
            r += 1
    wb.save(path)
